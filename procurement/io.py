"""Leitura de snapshots e contrato explícito de exportações CSV/XLSX."""
import csv
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
import hashlib
import io
import json
from pathlib import Path
import re

PR_FIELDS = ('rc','item_rc','data_rc','data_liberacao','solicitante','centro','grupo_compras','material','descricao','unidade','quantidade','status')
PO_FIELDS = ('pedido','item_pedido','rc','item_rc','data_pedido','fornecedor','centro','grupo_compras','material','descricao','unidade','quantidade','quantidade_recebida','preco_brl','unidade_preco','data_prevista','status')

def text(value):
    if value is None: return ''
    if isinstance(value, (datetime,date)): return value.date().isoformat() if isinstance(value,datetime) else value.isoformat()
    if isinstance(value,float) and value.is_integer(): return str(int(value))
    return str(value).strip()

def identifier(value, width, optional=False):
    value = text(value)
    if not value and optional: return ''
    if not re.fullmatch(r'\d{1,'+str(width)+r'}', value):
        raise ValueError(f'Identificador deve ter de 1 a {width} dígitos: {value!r}')
    return value.zfill(width)

def number(value, places, mapping, integer=False):
    if isinstance(value,bool): raise ValueError('Booleano não é quantidade/preço')
    if isinstance(value,(int,float,Decimal)):
        value = str(value)
    else:
        value = text(value)
        decimal, thousand = mapping['decimal'], mapping['milhar']
        whole = r'(?:\d{1,3}(?:'+re.escape(thousand)+r'\d{3})+|\d+)' if thousand else r'\d+'
        pattern = whole + (r'(?:'+re.escape(decimal)+r'\d{1,'+str(places)+r'})?' if places else '')
        if not re.fullmatch(pattern, value): raise ValueError('Número incompatível com o formato configurado: '+value)
        if thousand: value = value.replace(thousand,'')
        value = value.replace(decimal,'.')
    try: amount = Decimal(value)
    except InvalidOperation: raise ValueError('Número inválido') from None
    if not amount.is_finite() or amount < 0 or amount.as_tuple().exponent < -places:
        raise ValueError('Número negativo, não finito ou com casas demais')
    limit = 1000000 if places == 3 or integer else 10000000
    if amount > limit: raise ValueError('Número acima do limite demonstrativo')
    if integer and amount != amount.to_integral_value(): raise ValueError('Unidade de preço deve ser inteira')
    return int(amount * 10**places)

def input_date(value, mapping, optional=False):
    if not text(value) and optional: return ''
    if isinstance(value,(date,datetime)): return text(value)
    value = text(value)
    pattern = '%d/%m/%Y' if mapping['formato_data']=='br' else '%Y-%m-%d'
    check = r'\d{2}/\d{2}/\d{4}' if mapping['formato_data']=='br' else r'\d{4}-\d{2}-\d{2}'
    if not re.fullmatch(check,value): raise ValueError('Data incompatível com o formato: '+value)
    return datetime.strptime(value,pattern).date().isoformat()

def load_mapping(source):
    mapping = source if isinstance(source,dict) else json.loads(Path(source).read_text(encoding='utf-8-sig'))
    if mapping.get('formato_data') not in ('iso','br') or mapping.get('decimal') not in ('.',',') or mapping.get('milhar') not in ('','.',',') or mapping.get('decimal') == mapping.get('milhar') or mapping.get('separador') not in (';',',','\t'):
        raise ValueError('Formato de entrada inválido no mapeamento')
    for key,fields in [('requisicoes',PR_FIELDS),('pedidos',PO_FIELDS)]:
        spec = mapping.get(key,{})
        columns = spec.get('colunas',{})
        if set(columns) != set(fields) or any(not isinstance(v,str) or not v.strip() for v in columns.values()) or len(set(columns.values())) != len(fields):
            raise ValueError('Mapeamento incompleto ou repetido: '+key)
        name = spec.get('arquivo','')
        if not isinstance(name,str) or not name or Path(name).name != name or '/' in name or '\\' in name or Path(name).suffix.lower() not in ('.csv','.xlsx'):
            raise ValueError('Informe somente o nome do CSV/XLSX de entrada')
        if not isinstance(spec.get('aba'),str) or not spec['aba']: raise ValueError('Aba não informada')
    return mapping

def _read(path, spec, mapping, content):
    if path.suffix.lower()=='.csv':
        reader = csv.reader(io.StringIO(content.decode('utf-8-sig')),delimiter=mapping['separador'],strict=True)
        headers = next(reader,None)
        rows=[]
        while True:
            line=reader.line_num+1
            raw=next(reader,None)
            if raw is None: break
            if not any(text(v) for v in raw): continue
            if len(raw)!=len(headers or []): raise ValueError(f'{path.name}, linha {line}: quantidade de colunas diferente')
            rows.append((line,raw))
    else:
        from openpyxl import load_workbook
        workbook=load_workbook(io.BytesIO(content),data_only=False,read_only=True)
        try:
            if spec['aba'] not in workbook.sheetnames: raise ValueError('Aba não encontrada: '+spec['aba'])
            iterator=workbook[spec['aba']].iter_rows()
            first=next(iterator,None)
            headers=[text(c.value) for c in first] if first else None
            rows=[]
            for line,cells in enumerate(iterator,2):
                if not any(c.value is not None for c in cells): continue
                if any(c.data_type in ('f','e') for c in cells): raise ValueError(f'{path.name}, linha {line}: fórmula ou erro Excel; exporte valores conferidos')
                rows.append((line,[c.value for c in cells]))
        finally: workbook.close()
    if not headers or len(headers)!=len(set(headers)) or not set(spec['colunas'].values()).issubset(headers):
        raise ValueError(path.name+': cabeçalho ausente, repetido ou sem as colunas mapeadas')
    indices={field:headers.index(name) for field,name in spec['colunas'].items()}
    return [(line,{field:raw[index] for field,index in indices.items()}) for line,raw in rows]

def load_input(folder, mapping_source, as_of):
    cutoff=date.fromisoformat(as_of)
    mapping=load_mapping(mapping_source)
    base=Path(folder).resolve()
    loaded={'requisicoes':[],'pedidos':[],'fontes':{},'mapeamento':mapping}
    for key in ('requisicoes','pedidos'):
        spec=mapping[key]; path=base/spec['arquivo']; content=path.read_bytes()
        loaded['fontes'][path.name]={'sha256':hashlib.sha256(content).hexdigest(),'bytes':len(content),'base':key}
        for line,raw in _read(path,spec,mapping,content):
            try:
                row={f:text(v) for f,v in raw.items()}
                for f in ('rc','item_rc'):
                    row[f]=identifier(raw[f],5 if f=='item_rc' else 10,optional=key=='pedidos')
                if key=='pedidos':
                    row['pedido']=identifier(raw['pedido'],10); row['item_pedido']=identifier(raw['item_pedido'],5)
                    if bool(row['rc'])!=bool(row['item_rc']): raise ValueError('Referência RC/item precisa estar completa ou inteiramente vazia')
                for f in (('data_rc','data_liberacao') if key=='requisicoes' else ('data_pedido','data_prevista')):
                    row[f]=input_date(raw[f],mapping,optional=f=='data_liberacao')
                creation=row['data_rc' if key=='requisicoes' else 'data_pedido']
                if date.fromisoformat(creation)>cutoff: raise ValueError('Criação posterior à data de corte')
                if key=='requisicoes' and row['data_liberacao']:
                    if row['data_liberacao']<creation or row['data_liberacao']>as_of: raise ValueError('Data de liberação fora do intervalo criação/corte')
                allowed=('LIBERADA','AGUARDANDO','EXCLUIDA','REJEITADA') if key=='requisicoes' else ('ABERTO','CONCLUIDO','EXCLUIDO')
                if row['status'] not in allowed: raise ValueError('Status precisa seguir o contrato: '+', '.join(allowed))
                if key=='requisicoes' and row['status']=='LIBERADA' and not row['data_liberacao']: raise ValueError('RC liberada exige data de liberação conferida')
                for f in ('centro','grupo_compras','material','descricao','unidade', 'solicitante' if key=='requisicoes' else 'fornecedor'):
                    if not row[f] or '\x00' in row[f]: raise ValueError('Campo vazio ou NUL: '+f)
                row['quantidade_milesimos']=number(raw['quantidade'],3,mapping); del row['quantidade']
                if row['quantidade_milesimos']==0: raise ValueError('Quantidade do item precisa ser maior que zero')
                if key=='pedidos':
                    row['recebido_milesimos']=number(raw['quantidade_recebida'],3,mapping); del row['quantidade_recebida']
                    row['preco_centavos']=number(raw['preco_brl'],2,mapping); del row['preco_brl']
                    row['unidade_preco']=number(raw['unidade_preco'],0,mapping,integer=True)
                    if not row['unidade_preco']: raise ValueError('Unidade de preço precisa ser maior que zero')
                    if row['data_prevista']<creation: raise ValueError('Entrega prevista anterior à criação do pedido')
                row.update(arquivo=path.name,linha=line,campos_origem={f:text(v) for f,v in raw.items()})
                loaded[key].append(row)
            except (ValueError,InvalidOperation) as exc:
                raise ValueError(f'{path.name}, linha {line}: {exc}') from exc
    return loaded
